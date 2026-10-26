import sys
# Erzeugt die Vorlage PDF2E.xlsx (Blätter, Feldnamen, Formeln, Auswahllisten).
# Aufruf: python build_xlsx.py <Zieldatei.xlsx>   (benötigt openpyxl)
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
out=sys.argv[1]
wb=Workbook()
H=Font(bold=True,color="FFFFFF"); HF=PatternFill("solid",fgColor="7A4A12")
SEC=Font(bold=True,size=12,color="7A4A12")
INP=PatternFill("solid",fgColor="FFF8E1"); CALC=PatternFill("solid",fgColor="E8F0E8"); PDF=PatternFill("solid",fgColor="EEEEEE")
def header(ws,cols,widths):
    for i,(c,w) in enumerate(zip(cols,widths),1):
        x=ws.cell(1,i,c); x.font=H; x.fill=HF
        ws.column_dimensions[x.column_letter].width=w
    ws.freeze_panes="A2"
def name(n,ref): wb.defined_names[n]=DefinedName(n,attr_text=ref)

# Start
ws=wb.active; ws.title="Start"
ws.column_dimensions['A'].width=4; ws.column_dimensions['B'].width=26; ws.column_dimensions['C'].width=100
lines=[("PDF2E",Font(bold=True,size=22,color="7A4A12")),
("PDF-Rechnungen als E-Rechnung (ZUGFeRD)  |  Version 1.2.0",Font(size=11,color="555555")),
("",None),
("Ablauf",SEC),
("1. Makro 'RechnungEinlesen' starten und die BSI-PDF auswählen.",None),
("2. Blatt 'Rechnung' und 'Positionen' prüfen, gelb hinterlegte Felder bei Bedarf ergänzen (Zahlungskondition, Bankkonto).",None),
("3. Status im Blatt 'Rechnung' muss 'OK' zeigen. Bei Fehlern das Blatt 'Protokoll' ansehen; nach Korrekturen 'RechnungPruefen' erneut ausführen.",None),
("4. Makro 'ZugferdErzeugen' starten. Die ZUGFeRD-PDF landet im Ausgabeordner, das Original bleibt unverändert.",None),
("   Solange in 'Einstellungen' der Adapter STUB eingestellt ist, entsteht nur eine Kopie der PDF plus Datei *.daten.xml zur Kontrolle (keine E-Rechnung).",Font(bold=True,color="B00000")),
("",None),
("Einrichtung (einmalig)",SEC),
("- Blatt 'Einstellungen': Pfad zu pdftotext.exe, Eingangs- und Ausgabeordner.",None),
("- Blätter 'Brauerei', 'Bankkonten', 'Zahlungskonditionen' pflegen; Standard-Einträge mit 'x' markieren.",None),
("- Blatt 'Empfaenger': neue Kunden werden beim Einlesen automatisch angelegt. Nur abweichendes Konto, Kondition oder E-Mail eintragen (leer = Standard).",None),
("- Schaltflächen: einmalig das Makro 'PDF2E_Einrichten' ausführen (Alt+F8). Es legt die Schaltflächen hier und im Blatt 'Rechnung' an.",None),
("",None),
("Farben: gelb = Eingabe/änderbar, grün = berechnet, grau = Wert laut PDF (zum Abgleich).",Font(italic=True))]
for i,(t,f) in enumerate(lines,1):
    c=ws.cell(i,3,t); c.alignment=Alignment(wrap_text=True)
    if f: c.font=f

# Rechnung
ws=wb.create_sheet("Rechnung")
ws.column_dimensions['A'].width=30; ws.column_dimensions['B'].width=45; ws.column_dimensions['C'].width=55
rows=[
("Status",None),
("Prüfstatus","f_Status","calc",""),
("Quelldatei","f_Quelldatei","pdf",""),
("Ausgabedatei","f_Ausgabedatei","calc",""),
("Rechnung",None),
("Rechnungsnummer","f_RechNr","inp","BT-1"),
("Rechnungsdatum","f_RechDatum","inp","BT-2"),
("Rechnungsart","f_RechArt","inp","BT-3: 380 = Rechnung, 381 = Gutschrift"),
("Währung","f_Waehrung","inp","BT-5"),
("Lieferschein-Nr.","f_LSNr","inp","BT-16"),
("Leistungsdatum","f_LSDatum","inp","BT-72 (Lieferscheindatum)"),
("Zahlungskondition","f_Kondition","inp","Auswahl aus Blatt 'Zahlungskonditionen'"),
("Fälligkeitsdatum","f_Faellig","calc","BT-9 = Rechnungsdatum + Zahlungsziel"),
("Bankkonto","f_Konto","inp","Auswahl aus Blatt 'Bankkonten'"),
("Rechnungsempfänger",None),
("Kundennummer","f_KundenNr","inp","BT-46"),
("GLN Rechnungsempfänger","f_KGLN","inp","BT-46 mit Schema 0088 (z. B. RE-GLN bei REWE)"),
("Name","f_KName","inp","BT-44"),
("Straße","f_KStrasse","inp","BT-50"),
("PLZ","f_KPLZ","inp","BT-53"),
("Ort","f_KOrt","inp","BT-52"),
("Land","f_KLand","inp","BT-55"),
("E-Mail","f_KEmail","inp","BT-49, später auch Versandadresse"),
("USt-IdNr.","f_KUstId","inp","BT-48 (optional)"),
("Lieferort",None),
("Name","f_LName","inp","BT-70"),
("Zusatz/Straße","f_LZeile2","inp","BT-75"),
("PLZ","f_LPLZ","inp","BT-78"),
("Ort","f_LOrt","inp","BT-77"),
("Land","f_LLand","inp","BT-80"),
("GLN (ILN)","f_LGLN","inp","BT-71, Schema 0088"),
("Marktnummer","f_Marktnr","inp","nur Information"),
("Summen",None),
("Nettosumme (berechnet)","f_Netto","calc","Summe Positionen inkl. Pfand und Aktionen"),
("MwSt (berechnet)","f_MwSt","calc",""),
("Rechnungsbetrag (berechnet)","f_Brutto","calc",""),
("Warenwert laut PDF","f_PDF_Warenwert","pdf",""),
("Pfand laut PDF","f_PDF_Pfand","pdf",""),
("Zwischensumme laut PDF","f_PDF_Netto","pdf",""),
("MwSt laut PDF","f_PDF_MwSt","pdf",""),
("Rechnungsbetrag laut PDF","f_PDF_Brutto","pdf",""),
]
r=1; cells={}
for row in rows:
    if row[1] is None:
        r+=1 if r>1 else 0
        ws.cell(r,1,row[0]).font=SEC; r+=1; continue
    lab,nm,kind,hint=row
    ws.cell(r,1,lab); c=ws.cell(r,2); ws.cell(r,3,hint).font=Font(italic=True,color="777777")
    c.fill={"inp":INP,"calc":CALC,"pdf":PDF}[kind]
    name(nm,f"Rechnung!$B${r}"); cells[nm]=f"B{r}"; r+=1
for nm in ["f_RechDatum","f_LSDatum","f_Faellig"]: ws[cells[nm]].number_format="DD.MM.YYYY"
for nm in ["f_Netto","f_MwSt","f_Brutto","f_PDF_Warenwert","f_PDF_Pfand","f_PDF_Netto","f_PDF_MwSt","f_PDF_Brutto"]:
    ws[cells[nm]].number_format='#,##0.00'
for nm in ["f_RechNr","f_KundenNr","f_KGLN","f_KPLZ","f_LPLZ","f_LGLN","f_Marktnr","f_LSNr","f_RechArt"]: ws[cells[nm]].number_format='@'
ws[cells["f_RechArt"]]="380"; ws[cells["f_Waehrung"]]="EUR"; ws[cells["f_KLand"]]="DE"; ws[cells["f_LLand"]]="DE"
ws[cells["f_Faellig"]]='=IF(f_RechDatum="","",f_RechDatum+IFERROR(VLOOKUP(f_Kondition,Zahlungskonditionen!$A$2:$C$50,3,FALSE),0))'
ws[cells["f_Netto"]]="=SUM(Positionen!$G$2:$G$1000)"
ws[cells["f_MwSt"]]="=ROUND(SUMPRODUCT(Positionen!$G$2:$G$1000,Positionen!$H$2:$H$1000)/100,2)"
ws[cells["f_Brutto"]]="=f_Netto+f_MwSt"
ws[cells["f_Status"]].font=Font(bold=True)
dv=DataValidation(type="list",formula1="=Zahlungskonditionen!$A$2:$A$50",allow_blank=True); ws.add_data_validation(dv); dv.add(cells["f_Kondition"])
dv=DataValidation(type="list",formula1="=Bankkonten!$A$2:$A$20",allow_blank=True); ws.add_data_validation(dv); dv.add(cells["f_Konto"])

# Positionen
ws=wb.create_sheet("Positionen")
header(ws,["Typ","Art.-Nr.","Bezeichnung","Menge","Einheit","Einzelpreis","Betrag netto","USt-Satz %","USt-Kat.","Info (Hl./Fl.)"],[10,10,40,9,9,12,13,11,9,14])
for col in "FG": 
    for i in range(2,301): ws[f"{col}{i}"].number_format='#,##0.00##' if col=="F" else '#,##0.00'
dv=DataValidation(type="list",formula1='"Ware,Pfand,Nachlass"',allow_blank=True); ws.add_data_validation(dv); dv.add("A2:A1000")

# Protokoll
ws=wb.create_sheet("Protokoll")
header(ws,["Zeit","Art","Meldung"],[18,10,110])
for i in range(2,501): ws[f"A{i}"].number_format="DD.MM.YYYY hh:mm:ss"

# Brauerei
ws=wb.create_sheet("Brauerei")
header(ws,["Angabe","Wert","Hinweis"],[22,45,45])
data=[("Name","Konrad Krug Brauerei und Tanzsaal GmbH","BT-27"),("Straße","Breitenlesau 1b","BT-35"),("PLZ","91344","BT-38"),("Ort","Waischenfeld","BT-37"),("Land","DE","BT-40"),
("USt-IdNr","DE132316559","BT-31, von Frank bestätigt"),("Steuernummer","208/130/70031","BT-32"),("Registernummer","HRB 1273, Amtsgericht Bayreuth","BT-30"),
("Geschäftsführer","Konrad Krug",""),("Kontakt Name","","BT-41, bitte ergänzen"),("Telefon","09202/535","BT-42"),("E-Mail","info@krug-braeu.de","BT-43")]
for i,(a,b,c) in enumerate(data,2):
    ws.cell(i,1,a); x=ws.cell(i,2,b); x.fill=INP; x.number_format='@'; ws.cell(i,3,c).font=Font(italic=True,color="777777")

# Bankkonten
ws=wb.create_sheet("Bankkonten")
header(ws,["Bezeichnung","Kontoinhaber","IBAN","BIC","Standard (x)"],[24,40,28,16,13])
for i,row in enumerate([("Sparkasse Bayreuth","Konrad Krug Brauerei und Tanzsaal GmbH","DE58773501100000311274","BYLADEM1SBT","x"),
                        ("Raiffeisenbank Waischenfeld","Konrad Krug Brauerei und Tanzsaal GmbH","DE49773657920000212997","GENODEF1HWA","")],2):
    for j,v in enumerate(row,1): ws.cell(i,j,v).fill=INP

# Zahlungskonditionen
ws=wb.create_sheet("Zahlungskonditionen")
header(ws,["Bezeichnung","Text auf der E-Rechnung (BT-20)","Zahlungsziel Tage","Standard (x)"],[24,60,17,13])
for i,row in enumerate([("Sofort","Zahlbar sofort ohne Abzug.",0,"x"),("14 Tage netto","Zahlbar innerhalb von 14 Tagen ohne Abzug.",14,""),("30 Tage netto","Zahlbar innerhalb von 30 Tagen ohne Abzug.",30,"")],2):
    for j,v in enumerate(row,1): ws.cell(i,j,v).fill=INP

# Empfaenger
ws=wb.create_sheet("Empfaenger")
header(ws,["Kundennr.","Name","E-Mail Rechnung","USt-IdNr.","Vorgabe Bankkonto","Vorgabe Zahlungskondition","Bemerkung"],[11,36,30,16,26,26,30])
for j,v in enumerate(["11770","Markgrafen Getränkevertrieb GmbH","gam406@markgrafen.com","","","","Beispiel aus Muster RE 2603976; E-Mail bitte bestätigen"],1):
    c=ws.cell(2,j,v); c.fill=INP
ws["A2"].number_format='@'
dv=DataValidation(type="list",formula1="=Bankkonten!$A$2:$A$20",allow_blank=True); ws.add_data_validation(dv); dv.add("E2:E500")
dv=DataValidation(type="list",formula1="=Zahlungskonditionen!$A$2:$A$50",allow_blank=True); ws.add_data_validation(dv); dv.add("F2:F500")

# Einstellungen
ws=wb.create_sheet("Einstellungen")
header(ws,["Einstellung","Wert","Hinweis"],[22,50,60])
for i,(a,b,c) in enumerate([("pdftotext-Pfad",r"C:\Tools\xpdf\pdftotext.exe","Xpdf command line tools, 64-Bit-Version"),
("Eingangsordner",r"C:\Rechnungen\BSI","Ordner mit den BSI-PDFs"),
("Ausgabeordner",r"C:\Rechnungen\ZUGFeRD","Hier werden die ZUGFeRD-PDFs abgelegt"),
("Dateinamenszusatz","_ZUGFeRD","Dateiname = Rechnungsnummer + Zusatz + .pdf"),
("Layout","BSI","Einlese-Vorlage: BSI = Rechnungen aus der BSI-Software"),
("Layout-Variante","STANDARD","Abweichungen je Kunde innerhalb des BSI-Layouts; STANDARD = wie Muster Krug-Bräu"),
("Einheit Ware","C62","UN/ECE-Code; C62 = Stück (Kasten/Fass)"),
("Einheit Pfand","C62",""),
("Profil","EN16931","ZUGFeRD-Profil"),
("Quellcode-Ordner","","Nur Entwicklung: Zielordner für PDF2E_Exportieren (leer = Unterordner src neben der Mappe)"),
("Textdatei behalten","Nein","Ja = pdftotext-Ausgabe im TEMP-Ordner behalten (Fehlersuche), Pfad steht im Protokoll"),
("Entwickler","Frank Hentzschel","Wird im Info-Fenster angezeigt"),
("Firma","office & database development","Firmenbezeichnung für Info-Fenster und Copyright"),
("Kontakt","Frank.Hentzschel@gmx.de","z. B. E-Mail oder Webseite für Support"),
("Anwendungslogo","","Logo von PDF2E oben im Info-Fenster (BMP, JPG oder GIF, kein PNG), z. B. C:\\PDF2E\\pdf2e_logo.jpg"),
("Firmenlogo","","Logo des Entwicklers unten im Info-Fenster (BMP, JPG oder GIF, kein PNG)"),
("E-Rechnung-Adapter","STUB","STUB = Testlauf ohne Smart Tools (keine gültige E-Rechnung), SMARTTOOLS = echte Erzeugung")],2):
    ws.cell(i,1,a); x=ws.cell(i,2,b); x.fill=INP; ws.cell(i,3,c).font=Font(italic=True,color="777777")
ws=wb.create_sheet("Versionshistorie")
header(ws,["Version","Datum","Art","Änderung"],[10,12,16,100])
for i,row in enumerate([("1.2.0","26.10.2026","Neue Funktion","Neue Rechnungsempfänger werden automatisch im Blatt Empfaenger angelegt; Fehlerbehebung: umgebrochene Leergutzeilen der echten Windows-Ausgabe (Pfandteil vor der Artikelzeile) werden spaltengenau zusammengeführt"),
("1.1.0","26.10.2026","Neue Funktion","Info-Fenster: Anwendungslogo oben, Firmenlogo des Entwicklers unten (Einstellungen \"Anwendungslogo\" und \"Firmenlogo\")"),
("1.0.3","26.10.2026","Fehlerbehebung","Fehler beim Laden des Logos werden im Protokoll gemeldet"),
("1.0.2","26.10.2026","Fehlerbehebung","Schaltflächen als Formen statt Formularsteuerelemente (Anzeigefehler); Versionszeile im Blatt Start wird automatisch nachgeführt"),
("1.0.1","26.10.2026","Fehlerbehebung","Umgebrochene Leergutzeilen (pdftotext unter Windows) werden zusammengeführt; Option \"Textdatei behalten\""),
("1.0.0","26.10.2026","Erstversion","Einlesen BSI-PDF, Prüfung, Stub-Adapter, Splash- und Info-Fenster, Schaltflächen")],2):
    for j,v in enumerate(row,1): ws.cell(i,j,v)
ws["A2"].number_format='@'
wb.save(out)
